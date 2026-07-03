"""App-wide results export: /api/export/results.xlsx.

Exports the STORED detection answers (latest analysis_results per target); it
never recomputes. Uses the real Phase 1 sample fixture so the asserted answer
(user A -> home Hamra, High) is the frozen suite's own ground truth."""

import io

import pytest
from openpyxl import load_workbook

from app import create_app
from app.analysis.export import AREAS_HEADERS, CONTACTS_HEADERS, HEADERS, XLSX_MIMETYPE


@pytest.fixture()
def client(app_data_dir):
    app = create_app()
    app.config.update(TESTING=True)
    c = app.test_client()
    c.post("/api/setup", json={"username": "nadim", "password": "hunter2"})
    return c


@pytest.fixture()
def two_targets(client, sample_xlsx):
    """Target A analyzed, target B never analyzed. Returns (id_a, id_b)."""
    with open(sample_xlsx, "rb") as fh:
        data = {"name": "d", "file": (io.BytesIO(fh.read()), "sample.xlsx")}
    ds = client.post("/api/imports", data=data,
                     content_type="multipart/form-data").get_json()["data"]["dataset_id"]
    id_a = client.post("/api/targets", json={"name": "A", "dataset_id": ds,
                       "phone_number": "01000001"}).get_json()["data"]["id"]
    id_b = client.post("/api/targets", json={"name": "B", "dataset_id": ds,
                       "phone_number": "01000002"}).get_json()["data"]["id"]
    assert client.post(f"/api/targets/{id_a}/analyze").status_code == 200
    return id_a, id_b


def _sheet(res):
    return load_workbook(io.BytesIO(res.data)).active


def test_export_requires_auth(app_data_dir):
    app = create_app()
    app.config.update(TESTING=True)
    assert app.test_client().post("/api/export/results.xlsx").status_code == 401


def test_export_rejects_get(client, two_targets):
    """Export writes an EXPORT audit row, so it is POST-only: a GET (and the
    HEAD Flask would auto-add to a GET) must not reach the handler and log a
    phantom export where no bytes actually left the store."""
    assert client.get("/api/export/results.xlsx").status_code == 405


def test_export_all_targets(client, two_targets):
    res = client.post("/api/export/results.xlsx")
    assert res.status_code == 200
    assert res.mimetype == XLSX_MIMETYPE
    assert "attachment" in res.headers["Content-Disposition"]

    ws = _sheet(res)
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == HEADERS
    assert len(rows) == 3                       # header + both targets
    analyzed = dict(zip(HEADERS, rows[1]))
    assert analyzed["Target"] == "A"
    assert analyzed["Home area"] == "Hamra"
    assert analyzed["Home confidence"] == "High"
    assert analyzed["Frequent areas"].startswith("Hamra (")
    not_analyzed = dict(zip(HEADERS, rows[2]))
    assert not_analyzed["Target"] == "B"
    assert not_analyzed["Analyzed at"] == "not analyzed"
    assert not_analyzed["Home area"] in (None, "")


def test_export_has_no_confidence_on_tallies(client, two_targets):
    """Frequent areas/contacts are counts, not gated verdicts — no confidence
    columns anywhere in the workbook except Home/Work on the Summary sheet."""
    assert "Frequent areas confidence" not in HEADERS
    for headers in (AREAS_HEADERS, CONTACTS_HEADERS):
        assert not any("confidence" in h.lower() for h in headers)


def test_export_detail_sheets(client, two_targets):
    res = client.post("/api/export/results.xlsx")
    wb = load_workbook(io.BytesIO(res.data))
    assert wb.sheetnames == ["Summary", "Frequent areas", "Frequent contacts"]

    areas = list(wb["Frequent areas"].iter_rows(values_only=True))
    assert list(areas[0]) == AREAS_HEADERS
    # Target A (analyzed) contributes ranked area rows; Hamra is its top area.
    a_rows = [r for r in areas[1:] if r[1] == "A"]
    assert a_rows and a_rows[0][3] == 1 and a_rows[0][4] == "Hamra"
    # Target B (not analyzed) contributes none.
    assert not any(r[1] == "B" for r in areas[1:])

    contacts = list(wb["Frequent contacts"].iter_rows(values_only=True))
    assert list(contacts[0]) == CONTACTS_HEADERS
    assert all(r[1] == "A" for r in contacts[1:])


def test_export_single_target_scope(client, two_targets):
    id_a, _ = two_targets
    ws = _sheet(client.post(f"/api/export/results.xlsx?target_id={id_a}"))
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 2                       # header + just target A
    assert rows[1][1] == "A"


def test_export_unknown_target_404(client, two_targets):
    assert client.post("/api/export/results.xlsx?target_id=999").status_code == 404


def test_export_bad_target_id_400(client, two_targets):
    assert client.post("/api/export/results.xlsx?target_id=abc").status_code == 400


def test_export_writes_audit_row(client, two_targets):
    client.post("/api/export/results.xlsx")
    # The audit trail is the point of auditing EXPORT: data leaving the
    # encrypted store. Read it back through the app's own DB.
    from app.auth import account
    from app.infra import storage
    key = account.authenticate("nadim", "hunter2")
    conn = storage.connect(key)
    row = conn.execute(
        "SELECT username, detail FROM audit_log WHERE event_type = 'EXPORT'").fetchone()
    conn.close()
    assert row is not None
    assert row["username"] == "nadim"
    assert '"targets": 2' in row["detail"]
